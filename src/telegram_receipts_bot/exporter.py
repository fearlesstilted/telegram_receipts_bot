from __future__ import annotations

import json
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any

from .categories import CATEGORIES
from .state import current_month_key

REGISTER_FILENAME = "rozliczenie_paragonow.xlsx"
CATEGORY_COLUMN_START = 16

EXPORT_HEADERS = [
    "Data dok.",
    "Typ",
    "Wyjazd",
    "Data wyjazdu",
    "Nr FV",
    "Nr paragonu",
    "Kwota",
    "Kategoria",
    "Sprzedawca",
    "Adres",
    "NIP",
    "Towar",
    "Uwagi",
    "Budżet pozostały",
    "Zdjęcie/link",
    "Dodano",
]

RAW_HEADERS = [
    "created_at",
    "items_json",
    "ocr_text",
    "receipt_url",
]

REGISTER_HEADERS = [
    "Miesiąc",
    "Data dok.",
    "Typ",
    "Wyjazd",
    "Data wyjazdu",
    "Nr FV",
    "Nr paragonu",
    "Kwota",
    "Kategoria",
    "Sprzedawca",
    "Adres",
    "NIP",
    "Towar",
    "Uwagi",
    "Pozostało po dokumencie",
    *CATEGORIES,
    "Zdjęcie/link",
    "Dodano",
]

MONTH_HEADERS = [
    "Miesiąc",
    "Zaliczka",
    "Suma paragonów",
    "Pozostało",
    "Liczba dokumentów",
]


def generate_export(data_dir: Path, month_key: str | None = None) -> Path:
    from openpyxl import Workbook

    month_key = month_key or current_month_key()
    rows = _read_rows(data_dir / "rows.jsonl")
    rows = [row for row in rows if _row_month(row) == month_key]

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Podsumowanie"
    sheet = workbook.create_sheet("Paragony")
    raw_sheet = workbook.create_sheet("OCR raw")
    categories_sheet = workbook.create_sheet("Kategorie")

    _write_main_sheet(sheet, rows)
    _write_summary_sheet(summary_sheet, month_key, rows)
    _write_raw_sheet(raw_sheet, rows)
    _write_categories_sheet(categories_sheet)
    _apply_category_validation(sheet, categories_sheet)
    categories_sheet.sheet_state = "hidden"

    target = data_dir / f"export_{month_key}.xlsx"
    workbook.save(target)
    return target


def generate_register(data_dir: Path) -> Path:
    from openpyxl import Workbook

    rows = _read_rows(data_dir / "rows.jsonl")
    months = sorted({_row_month(row) for row in rows} | {current_month_key()})
    target = data_dir / REGISTER_FILENAME
    advances = _read_existing_advances(target)

    workbook = Workbook()
    dashboard_sheet = workbook.active
    dashboard_sheet.title = "Panel"
    months_sheet = workbook.create_sheet("Miesiące")
    receipts_sheet = workbook.create_sheet("Paragony")
    raw_sheet = workbook.create_sheet("OCR raw")
    categories_sheet = workbook.create_sheet("Kategorie")

    _write_dashboard_sheet(dashboard_sheet, months, rows)
    _write_months_sheet(months_sheet, months, advances)
    _write_register_sheet(receipts_sheet, rows)
    _write_raw_sheet(raw_sheet, rows)
    _write_categories_sheet(categories_sheet)
    _apply_category_validation(receipts_sheet, categories_sheet, "I2:I5000")
    categories_sheet.sheet_state = "hidden"
    _apply_workbook_theme(workbook)

    workbook.save(target)
    return target


def _read_rows(rows_file: Path) -> list[dict[str, Any]]:
    if not rows_file.exists():
        return []

    rows: list[dict[str, Any]] = []
    with rows_file.open(encoding="utf-8") as file_obj:
        for line in file_obj:
            line = line.strip()
            if not line:
                continue
            rows.append(json.loads(line))
    return rows


def _read_existing_advances(workbook_path: Path) -> dict[str, float | str]:
    if not workbook_path.exists():
        return {}

    try:
        from openpyxl import load_workbook

        workbook = load_workbook(workbook_path, data_only=False)
    except Exception:
        return {}

    if "Miesiące" not in workbook.sheetnames:
        return {}

    advances: dict[str, float | str] = {}
    sheet = workbook["Miesiące"]
    for row_idx in range(2, sheet.max_row + 1):
        month = str(sheet.cell(row=row_idx, column=1).value or "").strip()
        if not month:
            continue
        value = sheet.cell(row=row_idx, column=2).value
        advances[month] = value if value is not None else ""
    return advances


def _row_month(row: dict[str, Any]) -> str:
    date_value = str(row.get("data_dokumentu") or row.get("created_at") or "")
    if len(date_value) >= 7:
        return date_value[:7]
    return current_month_key()


def _write_main_sheet(sheet, rows: list[dict[str, Any]]) -> None:
    from openpyxl.styles import Font

    sheet.append(EXPORT_HEADERS)
    for row in rows:
        sheet.append(
            [
                row.get("data_dokumentu", ""),
                row.get("typ_dokumentu", ""),
                row.get("wyjazd_nazwa", ""),
                row.get("wyjazd_data", ""),
                row.get("nr_fv", ""),
                row.get("nr_paragonu", ""),
                _to_float(row.get("kwota")),
                row.get("kategoria", ""),
                row.get("sprzedawca", ""),
                row.get("adres", ""),
                row.get("nip", ""),
                row.get("towar", ""),
                row.get("uwagi", ""),
                _to_float(row.get("budget_remaining")),
                row.get("receipt_url", ""),
                _format_created_at(row.get("created_at", "")),
            ]
        )

    total_row = sheet.max_row + 1
    sheet.cell(row=total_row, column=6, value="Razem")
    if total_row > 2:
        sheet.cell(row=total_row, column=7, value=f"=SUM(G2:G{total_row - 1})")
    else:
        sheet.cell(row=total_row, column=7, value=0)
    for cell in sheet[total_row]:
        cell.font = Font(bold=True)

    _style_table(sheet, money_columns=(7, 14))


def _write_register_sheet(sheet, rows: list[dict[str, Any]]) -> None:
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheet.append(REGISTER_HEADERS)
    for row_index, row in enumerate(rows, start=2):
        month_key = _row_month(row)
        category_amounts = _category_amount_formulas(row_index)
        sheet.append(
            [
                month_key,
                row.get("data_dokumentu", ""),
                row.get("typ_dokumentu", ""),
                row.get("wyjazd_nazwa", ""),
                row.get("wyjazd_data", ""),
                row.get("nr_fv", ""),
                row.get("nr_paragonu", ""),
                _to_float(row.get("kwota")),
                row.get("kategoria", ""),
                row.get("sprzedawca", ""),
                row.get("adres", ""),
                row.get("nip", ""),
                row.get("towar", ""),
                row.get("uwagi", ""),
                (
                    f'=IF(A{row_index}="","",'
                    f'IFERROR(INDEX(Miesiące!$B:$B,MATCH(A{row_index},Miesiące!$A:$A,0)),0)'
                    f'-SUMIFS($H$2:H{row_index},$A$2:A{row_index},A{row_index}))'
                ),
                *category_amounts,
                row.get("receipt_url", ""),
                _format_created_at(row.get("created_at", "")),
            ]
        )

    total_row = sheet.max_row + 1
    sheet.cell(row=total_row, column=7, value="Razem")
    if total_row > 2:
        sheet.cell(row=total_row, column=8, value=f"=SUM(H2:H{total_row - 1})")
        for offset, _category in enumerate(CATEGORIES):
            col_idx = CATEGORY_COLUMN_START + offset
            column_letter = get_column_letter(col_idx)
            sheet.cell(
                row=total_row,
                column=col_idx,
                value=f"=SUM({column_letter}2:{column_letter}{total_row - 1})",
            )
    else:
        sheet.cell(row=total_row, column=8, value=0)
        for offset, _category in enumerate(CATEGORIES):
            sheet.cell(row=total_row, column=CATEGORY_COLUMN_START + offset, value=0)
    for cell in sheet[total_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E2F0D9")

    category_money_columns = tuple(
        range(CATEGORY_COLUMN_START, CATEGORY_COLUMN_START + len(CATEGORIES))
    )
    _style_table(sheet, money_columns=(8, 15, *category_money_columns))
    _add_excel_table(sheet, "ParagonyTabela")
    _apply_workbook_sheet_layout(sheet)


def _write_months_sheet(sheet, months: list[str], advances: dict[str, float | str]) -> None:
    from openpyxl.styles import Font, PatternFill

    sheet.append(MONTH_HEADERS)
    for row_idx, month_key in enumerate(months, start=2):
        advance = advances.get(month_key, "")
        sheet.append(
            [
                month_key,
                advance,
                f'=SUMIF(Paragony!$A:$A,A{row_idx},Paragony!$H:$H)',
                f'=B{row_idx}-C{row_idx}',
                f'=COUNTIF(Paragony!$A:$A,A{row_idx})',
            ]
        )

    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="2E75B6")
        cell.font = Font(color="FFFFFF", bold=True)
    for row_idx in range(2, sheet.max_row + 1):
        for col_idx in (2, 3, 4):
            sheet.cell(row=row_idx, column=col_idx).number_format = '#,##0.00'
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_letter, width in {"A": 12, "B": 14, "C": 18, "D": 14, "E": 18}.items():
        sheet.column_dimensions[column_letter].width = width
    _add_excel_table(sheet, "MiesiaceTabela")
    _apply_workbook_sheet_layout(sheet)


def _write_dashboard_sheet(sheet, months: list[str], rows: list[dict[str, Any]]) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    current_month = current_month_key()
    month = current_month if current_month in months else (months[-1] if months else current_month)
    category_start_row = 12
    category_end_row = category_start_row + len(CATEGORIES) - 1

    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A12"
    sheet.sheet_properties.tabColor = "2F5597"

    dark_fill = PatternFill("solid", fgColor="2F5597")
    section_fill = PatternFill("solid", fgColor="D9EAF7")
    muted_fill = PatternFill("solid", fgColor="F4F7FA")
    good_fill = PatternFill("solid", fgColor="E2F0D9")
    warning_fill = PatternFill("solid", fgColor="FCE4D6")
    border = Border(
        left=Side(style="thin", color="D9E2F3"),
        right=Side(style="thin", color="D9E2F3"),
        top=Side(style="thin", color="D9E2F3"),
        bottom=Side(style="thin", color="D9E2F3"),
    )

    sheet.merge_cells("A1:F1")
    sheet["A1"] = "Rozliczenie dokumentów"
    sheet["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    sheet["A1"].fill = dark_fill
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 26

    summary_cells = [
        ("A3", "Aktywny miesiąc", "B3", month),
        ("C3", "Liczba dokumentów", "D3", f'=IFERROR(COUNTIF(Paragony!$A:$A,$B$3),0)'),
        ("E3", "Pozostało", "F3", "=B4-D4"),
        ("A4", "Zaliczka", "B4", f'=IFERROR(INDEX(Miesiące!$B:$B,MATCH($B$3,Miesiące!$A:$A,0)),0)'),
        ("C4", "Suma dokumentów", "D4", f'=IFERROR(SUMIF(Paragony!$A:$A,$B$3,Paragony!$H:$H),0)'),
        ("E4", "Ostatni zapis", "F4", '=IFERROR(MAX(Paragony!$B:$B),"")'),
    ]
    for label_cell, label, value_cell, value in summary_cells:
        sheet[label_cell] = label
        sheet[label_cell].font = Font(size=9, bold=True, color="44546A")
        sheet[label_cell].fill = section_fill
        sheet[label_cell].border = border
        sheet[value_cell] = value
        sheet[value_cell].font = Font(size=10, bold=True, color="1F1F1F")
        sheet[value_cell].fill = muted_fill
        sheet[value_cell].border = border
        sheet[value_cell].alignment = Alignment(horizontal="right", vertical="center")
    for cell_ref in ("B4", "D4", "F3"):
        sheet[cell_ref].number_format = '#,##0.00 "PLN"'
    sheet["F3"].fill = warning_fill

    sheet["A6"] = "Instrukcja"
    sheet["A6"].font = Font(bold=True, color="44546A")
    sheet["A6"].fill = section_fill
    sheet["A6"].border = border
    sheet.merge_cells("B6:F6")
    sheet["B6"] = "Wpisz zaliczkę w arkuszu Miesiące. Bot dopisuje dokumenty do arkusza Paragony. Kategorię można zmienić w kolumnie I."
    sheet["B6"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet["B6"].fill = muted_fill
    sheet["B6"].border = border

    sheet["A11"] = "Kategoria"
    sheet["B11"] = "Suma"
    sheet["C11"] = "Liczba"
    for cell in sheet[11][:3]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = dark_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

    for row_idx, category in enumerate(CATEGORIES, start=category_start_row):
        sheet.cell(row=row_idx, column=1, value=category)
        sheet.cell(
            row=row_idx,
            column=2,
            value=f'=IFERROR(SUMIFS(Paragony!$H:$H,Paragony!$A:$A,$B$3,Paragony!$I:$I,A{row_idx}),0)',
        )
        sheet.cell(
            row=row_idx,
            column=3,
            value=f'=IFERROR(COUNTIFS(Paragony!$A:$A,$B$3,Paragony!$I:$I,A{row_idx}),0)',
        )
        sheet.cell(row=row_idx, column=2).number_format = '#,##0.00'
        for col_idx in range(1, 4):
            cell = sheet.cell(row=row_idx, column=col_idx)
            cell.fill = muted_fill if row_idx % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")
            cell.border = border
            cell.alignment = Alignment(vertical="center")
        sheet.cell(row=row_idx, column=2).alignment = Alignment(horizontal="right")
        sheet.cell(row=row_idx, column=3).alignment = Alignment(horizontal="right")

    total_row = category_end_row + 1
    sheet.cell(row=total_row, column=1, value="Razem")
    sheet.cell(row=total_row, column=2, value=f"=SUM(B{category_start_row}:B{category_end_row})")
    sheet.cell(row=total_row, column=3, value=f"=SUM(C{category_start_row}:C{category_end_row})")
    for cell in sheet[total_row]:
        cell.font = Font(bold=True)
        cell.fill = good_fill
        cell.border = border

    sheet["D11"] = "Ostatnie dokumenty"
    sheet["D11"].font = Font(color="FFFFFF", bold=True)
    sheet["D11"].fill = dark_fill
    sheet["D11"].border = border
    sheet.merge_cells("D11:F11")
    sheet["D12"] = "Data"
    sheet["E12"] = "Kwota"
    sheet["F12"] = "Sprzedawca"
    for cell in sheet[12][3:6]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = dark_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center")
    recent_rows = rows[-5:][::-1]
    for row_idx, row in enumerate(recent_rows, start=13):
        sheet.cell(row=row_idx, column=4, value=row.get("data_dokumentu") or row.get("created_at", "")[:10])
        sheet.cell(row=row_idx, column=5, value=_to_float(row.get("kwota")))
        sheet.cell(row=row_idx, column=6, value=row.get("sprzedawca") or row.get("towar", ""))
        sheet.cell(row=row_idx, column=5).number_format = '#,##0.00'
        for col_idx in range(4, 7):
            cell = sheet.cell(row=row_idx, column=col_idx)
            cell.fill = muted_fill if row_idx % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=col_idx == 6)

    for column_letter, width in {"A": 26, "B": 14, "C": 10, "D": 13, "E": 12, "F": 34}.items():
        sheet.column_dimensions[column_letter].width = width
    for row_idx in range(2, 26):
        sheet.row_dimensions[row_idx].height = 21


def _write_summary_sheet(sheet, month_key: str, rows: list[dict[str, Any]]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    data_last_row = len(rows) + 1
    amount_range = f"Paragony!G2:G{data_last_row}"
    category_range = f"Paragony!H2:H{data_last_row}"

    sheet["A1"] = "Rozliczenie paragonów"
    sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="2E75B6")
    sheet.merge_cells("A1:D1")

    sheet["A3"] = "Miesiąc"
    sheet["B3"] = month_key
    sheet["A4"] = "Liczba dokumentów"
    sheet["B4"] = len(rows)
    sheet["A5"] = "Suma kosztów"
    sheet["B5"] = f"=SUM({amount_range})" if rows else 0
    sheet["B5"].number_format = '#,##0.00'
    sheet["A6"] = "Ostatni budżet pozostały"
    sheet["B6"] = _to_float(rows[-1].get("budget_remaining")) if rows else None
    sheet["B6"].number_format = '#,##0.00'

    sheet["A8"] = "Kategoria"
    sheet["B8"] = "Suma"
    sheet["C8"] = "Liczba"
    for cell in sheet[8]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="2E75B6")

    for row_idx, category in enumerate(CATEGORIES, start=9):
        sheet.cell(row=row_idx, column=1, value=category)
        if rows:
            sheet.cell(row=row_idx, column=2, value=f'=SUMIF({category_range},A{row_idx},{amount_range})')
            sheet.cell(row=row_idx, column=3, value=f'=COUNTIF({category_range},A{row_idx})')
        else:
            sheet.cell(row=row_idx, column=2, value=0)
            sheet.cell(row=row_idx, column=3, value=0)
        sheet.cell(row=row_idx, column=2).number_format = '#,##0.00'

    total_row = 9 + len(CATEGORIES)
    sheet.cell(row=total_row, column=1, value="Razem")
    sheet.cell(row=total_row, column=2, value=f"=SUM(B9:B{total_row - 1})")
    sheet.cell(row=total_row, column=3, value=f"=SUM(C9:C{total_row - 1})")
    for cell in sheet[total_row]:
        cell.font = Font(bold=True)

    for column_letter, width in {"A": 28, "B": 16, "C": 12, "D": 18}.items():
        sheet.column_dimensions[column_letter].width = width
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top")


def _write_raw_sheet(sheet, rows: list[dict[str, Any]]) -> None:
    sheet.append(RAW_HEADERS)
    for row in rows:
        sheet.append([row.get(header, "") for header in RAW_HEADERS])
    _style_table(sheet, money_columns=())
    _add_excel_table(sheet, "OcrRawTabela")
    _apply_workbook_sheet_layout(sheet)


def _write_categories_sheet(sheet) -> None:
    sheet.append(["Kategoria"])
    for category in CATEGORIES:
        sheet.append([category])
    _style_table(sheet, money_columns=())
    _apply_workbook_sheet_layout(sheet)


def _category_amount_formulas(row_index: int) -> list[str]:
    formulas = []
    for category in CATEGORIES:
        formulas.append(f'=IF($I{row_index}="{category}",$H{row_index},0)')
    return formulas


def _apply_category_validation(sheet, categories_sheet, cell_range: str = "H2:H1000") -> None:
    from openpyxl.worksheet.datavalidation import DataValidation

    last_category_row = len(CATEGORIES) + 1
    validation = DataValidation(
        type="list",
        formula1=f"=Kategorie!$A$2:$A${last_category_row}",
        allow_blank=True,
    )
    validation.error = "Wybierz kategorię z listy."
    validation.errorTitle = "Nieznana kategoria"
    validation.prompt = "Wybierz kategorię kosztu."
    validation.promptTitle = "Kategoria"
    sheet.add_data_validation(validation)
    validation.add(cell_range)


def _style_table(sheet, money_columns: tuple[int, ...]) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="2F5597")
    even_fill = PatternFill("solid", fgColor="F4F7FA")
    odd_fill = PatternFill("solid", fgColor="FFFFFF")
    thin = Side(style="thin", color="D9E2F3")

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx in range(2, sheet.max_row + 1):
        fill = even_fill if row_idx % 2 == 0 else odd_fill
        for cell in sheet[row_idx]:
            cell.fill = fill
            cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)

    for col_idx in money_columns:
        for row_idx in range(2, sheet.max_row + 1):
            sheet.cell(row=row_idx, column=col_idx).number_format = '#,##0.00'

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (11, 12, 14))
    for column_cells in sheet.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            max_length = max(max_length, len(str(cell.value or "")))
        sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 10), 55)
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 24


def _apply_workbook_theme(workbook) -> None:
    from openpyxl.styles import Font

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                cell.font = copy(cell.font)
                cell.font = Font(
                    name="Calibri",
                    size=10 if cell.font.sz is None else cell.font.sz,
                    bold=cell.font.bold,
                    italic=cell.font.italic,
                    color=cell.font.color,
                )
        sheet.sheet_view.showGridLines = False


def _add_excel_table(sheet, table_name: str) -> None:
    if sheet.max_row < 1 or sheet.max_column < 1:
        return
    from openpyxl.worksheet.table import Table, TableStyleInfo

    table = Table(displayName=table_name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def _apply_workbook_sheet_layout(sheet) -> None:
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.outlinePr.summaryBelow = False


def _to_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return round(float(str(value).replace(",", ".")), 2)
    except ValueError:
        return None


def _format_created_at(value: Any) -> str:
    raw = str(value or "")
    if not raw:
        return ""
    try:
        return datetime.fromisoformat(raw).strftime("%Y-%m-%d %H:%M")
    except ValueError:
        return raw
