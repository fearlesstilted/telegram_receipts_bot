import json

from openpyxl import load_workbook

from telegram_receipts_bot.categories import CATEGORIES
from telegram_receipts_bot.exporter import (
    CATEGORY_COLUMN_START,
    EXPORT_HEADERS,
    REGISTER_FILENAME,
    generate_export,
    generate_register,
)
from telegram_receipts_bot.models import ReceiptDraft
from telegram_receipts_bot.reset_excel import main as reset_excel_main
from telegram_receipts_bot.sinks import LocalJsonSink
from telegram_receipts_bot.state import current_month_key


def test_generate_export_creates_styled_workbook(tmp_path) -> None:
    rows_file = tmp_path / "rows.jsonl"
    rows_file.write_text(
        json.dumps(
            {
                "data_dokumentu": "2026-04-09",
                "typ_dokumentu": "paragon",
                "wyjazd_nazwa": "Brzeg",
                "wyjazd_data": "2026-04-09",
                "nr_fv": "",
                "nr_paragonu": "108592",
                "kwota": "127.00",
                "kategoria": "Wyżywienie",
                "sprzedawca": "CENTRUM USŁUG GASTRONOMICZNYCH MARIO",
                "adres": "ELEKTRYCZNA 8A, 49-300 BRZEG",
                "nip": "7471727805",
                "towar": "DANIE DNIA",
                "uwagi": "",
                "budget_remaining": "873.00",
                "items_json": "[]",
                "ocr_text": "SUMA: PLN 127.00",
                "receipt_url": "/tmp/r.jpg",
                "created_at": "2026-04-09T16:14:00+00:00",
            },
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )

    export_path = generate_export(tmp_path, "2026-04")

    workbook = load_workbook(export_path, data_only=False)
    sheet = workbook["Paragony"]
    assert [cell.value for cell in sheet[1]] == EXPORT_HEADERS
    assert sheet["B2"].value == "paragon"
    assert sheet["F2"].value == "108592"
    assert sheet["G2"].value == 127.00
    assert sheet["H2"].value == "Wyżywienie"
    assert sheet["J2"].value == "ELEKTRYCZNA 8A, 49-300 BRZEG"
    assert sheet["G3"].value == "=SUM(G2:G2)"
    assert workbook.sheetnames[:3] == ["Podsumowanie", "Paragony", "OCR raw"]
    summary = workbook["Podsumowanie"]
    assert summary["A1"].value == "Rozliczenie paragonów"
    assert summary["B3"].value == "2026-04"
    assert summary["B4"].value == 1
    assert summary["B5"].value == "=SUM(Paragony!G2:G2)"
    assert workbook["Kategorie"].sheet_state == "hidden"
    assert "OCR raw" in workbook.sheetnames


def test_local_sink_writes_rows_for_export(tmp_path) -> None:
    receipt_file = tmp_path / "source.jpg"
    receipt_file.write_bytes(b"demo")
    draft = ReceiptDraft.empty("abc", 123, str(receipt_file))
    draft.data_dokumentu = "2026-04-09"
    draft.typ_dokumentu = "paragon"
    draft.kwota = 42.50

    sink = LocalJsonSink(tmp_path)

    url = sink.save_receipt(draft, budget_remaining=957.50)

    assert url.endswith("source.jpg")
    rows = [
        json.loads(line)
        for line in (tmp_path / "rows.jsonl").read_text(encoding="utf-8").splitlines()
    ]
    assert rows[0]["data_dokumentu"] == "2026-04-09"
    assert rows[0]["typ_dokumentu"] == "paragon"
    assert rows[0]["kwota"] == "42.50"


def test_generate_register_preserves_monthly_advance(tmp_path) -> None:
    rows_file = tmp_path / "rows.jsonl"
    rows_file.write_text(
        json.dumps(
            {
                "data_dokumentu": "2026-04-09",
                "typ_dokumentu": "paragon",
                "nr_paragonu": "108592",
                "kwota": "127.00",
                "kategoria": "Wyżywienie",
                "sprzedawca": "CENTRUM USŁUG GASTRONOMICZNYCH MARIO",
                "created_at": "2026-04-09T16:14:00+00:00",
            },
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )

    register_path = generate_register(tmp_path)
    workbook = load_workbook(register_path)
    workbook["Miesiące"]["B2"] = 1000
    workbook.save(register_path)

    register_path = generate_register(tmp_path)
    workbook = load_workbook(register_path, data_only=False)

    assert register_path.name == REGISTER_FILENAME
    assert workbook.sheetnames[:4] == ["Panel", "Miesiące", "Paragony", "OCR raw"]
    panel = workbook["Panel"]
    assert panel["A1"].value == "Rozliczenie dokumentów"
    assert panel["B3"].value == current_month_key()
    assert panel["B4"].value == '=IFERROR(INDEX(Miesiące!$B:$B,MATCH($B$3,Miesiące!$A:$A,0)),0)'
    assert panel["D4"].value == '=IFERROR(SUMIF(Paragony!$A:$A,$B$3,Paragony!$H:$H),0)'
    assert panel["F3"].value == "=B4-D4"
    assert workbook["Miesiące"]["A2"].value == "2026-04"
    assert workbook["Miesiące"]["B2"].value == 1000
    assert workbook["Miesiące"]["C2"].value == "=SUMIF(Paragony!$A:$A,A2,Paragony!$H:$H)"
    assert workbook["Miesiące"]["D2"].value == "=B2-C2"
    assert workbook["Paragony"]["A2"].value == "2026-04"
    assert workbook["Paragony"]["C2"].value == "paragon"
    assert workbook["Paragony"]["H2"].value == 127.00
    assert workbook["Paragony"]["O2"].value.startswith('=IF(A2=""')
    assert "INDEX(Miesiące!$B:$B,MATCH" in workbook["Paragony"]["O2"].value
    assert workbook["Paragony"].cell(1, CATEGORY_COLUMN_START).value == CATEGORIES[0]
    assert workbook["Paragony"].cell(1, CATEGORY_COLUMN_START + 1).value == CATEGORIES[1]
    assert workbook["Paragony"].cell(2, CATEGORY_COLUMN_START).value == '=IF($I2="Wyżywienie",$H2,0)'
    assert workbook["Paragony"].cell(2, CATEGORY_COLUMN_START + 1).value == '=IF($I2="Paliwo",$H2,0)'
    total_row = workbook["Paragony"].max_row
    assert workbook["Paragony"].cell(total_row, CATEGORY_COLUMN_START).value.startswith("=SUM(")
    assert workbook["Kategorie"].sheet_state == "hidden"


def test_reset_excel_archives_rows_and_creates_empty_register(tmp_path, monkeypatch) -> None:
    rows_file = tmp_path / "rows.jsonl"
    rows_file.write_text(
        json.dumps({"data_dokumentu": "2026-04-09", "kwota": "2000000.00"}) + "\n",
        encoding="utf-8",
    )
    state_file = tmp_path / "state.json"
    state_file.write_text(
        json.dumps(
            {
                "monthly_budget": {"2026-04": {"total": 1000.0, "spent": 2000000.0}},
                "pending_drafts": {"abc": {}},
            }
        ),
        encoding="utf-8",
    )
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    monkeypatch.setenv("TELEGRAM_BOT_TOKEN", "token")
    monkeypatch.setenv("TELEGRAM_ALLOWED_CHAT_ID", "1")

    reset_excel_main()

    assert not rows_file.exists()
    assert (tmp_path / REGISTER_FILENAME).exists()
    archives = list((tmp_path / "archive").glob("*"))
    assert len(archives) == 1
    assert (archives[0] / "rows.jsonl").exists()
    state = json.loads(state_file.read_text(encoding="utf-8"))
    assert state["monthly_budget"]["2026-04"]["spent"] == 0.0
    assert state["pending_drafts"] == {}
